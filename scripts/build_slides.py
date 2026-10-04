#!/usr/bin/env python3
"""Build the public CAMG paper talk in the project's clean academic style.

The source figures come from the public project page. The generated PPTX keeps
text and geometry editable. render_slides_pdf.py exports the same layout to a
browser-friendly PDF without desktop automation.
"""
from __future__ import annotations

import argparse
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SW, SH = 10.0, 5.625
FONT = "Arial"
SERIF = "Times New Roman"
INK = "111111"
MID = "555555"
LIGHT = "8A8A8A"
LINE = "C9C9C9"
BLUE = "2B628F"
BLUE_LIGHT = "DCEAF3"
RED = "C8192E"
RED_LIGHT = "F8E5E8"
GREEN = "198754"
WHITE = "FFFFFF"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#").upper())


def set_run(run, size, color=INK, bold=False, italic=False, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)


def text(slide, content, x, y, w, h, size=18, color=INK, bold=False,
         italic=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
         font=FONT, margin=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for idx, line in enumerate(str(content).split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.line_spacing = 1.0
        p.space_after = Pt(0)
        for run in p.runs:
            set_run(run, size, color, bold, italic, font)
    return box


def rich_text(slide, runs, x, y, w, h, size=18, align=PP_ALIGN.LEFT,
              valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = 1.0
    for spec in runs:
        run = p.add_run()
        run.text = spec["text"]
        set_run(run, spec.get("size", size), spec.get("color", INK),
                spec.get("bold", False), spec.get("italic", False),
                spec.get("font", font))
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=INK, width=1.2, rounded=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(width)
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, width=1.0):
    shape = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def image(slide, path, x, y, w, h, contain=True):
    path = Path(path)
    im = Image.open(path)
    iw, ih = im.size
    target = w / h
    actual = iw / ih
    if contain:
        if actual > target:
            rw, rh = w, w / actual
            xx, yy = x, y + (h - rh) / 2
        else:
            rh, rw = h, h * actual
            xx, yy = x + (w - rw) / 2, y
        return slide.shapes.add_picture(str(path), Inches(xx), Inches(yy), Inches(rw), Inches(rh))
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def title(slide, heading, number=None):
    text(slide, heading, 0.38, 0.23, 8.7, 0.48, size=24, bold=False)
    text(slide, "JD.com", 8.82, 0.20, 0.76, 0.27, size=9, color=MID, bold=True, align=PP_ALIGN.RIGHT)
    line(slide, 0.38, 0.74, 9.62, 0.74, color=LINE, width=0.8)
    if number is not None:
        text(slide, f"{number:02d}", 9.38, 5.20, 0.25, 0.18, size=8, color=LIGHT, align=PP_ALIGN.RIGHT)


def caption(slide, label, copy, y=5.10):
    rich_text(slide, [{"text": label.upper() + "  ", "size": 8.5, "color": RED, "bold": True},
                      {"text": copy, "size": 10.0, "color": MID}],
              0.45, y, 8.85, 0.36)


def bullets(slide, items, x, y, w, h, size=16, gap=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + item
        p.space_after = Pt(gap); p.line_spacing = 1.02
        for run in p.runs: set_run(run, size, INK)
    return box


def make_chart_images(out_dir: Path):
    out_dir.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11})
    methods = ["Qwen3.5-4B", "Mem0", "Letta Code", "CompactionRL", "AgeMem", "CAMG-RL"]
    avg = [17.4, 19.7, 17.8, 48.0, 25.3, 54.4]
    fig, ax = plt.subplots(figsize=(9.0, 3.4), dpi=220)
    colors = ["#BFC9D3"] * 5 + ["#2B628F"]
    bars = ax.bar(methods, avg, color=colors, width=0.64)
    ax.set_ylim(0, 65); ax.set_ylabel("Success rate (%)")
    ax.spines[["top", "right"]].set_visible(False); ax.grid(axis="y", color="#E6E8EA", linewidth=.7); ax.set_axisbelow(True)
    ax.tick_params(axis="x", labelrotation=20, labelsize=9); ax.tick_params(axis="y", labelsize=9)
    for b, v in zip(bars, avg): ax.text(b.get_x()+b.get_width()/2, v+1.2, f"{v:.1f}", ha="center", fontsize=9, fontweight="bold" if v == 54.4 else "normal")
    fig.tight_layout(); fig.savefig(out_dir / "main-result.png", transparent=False, facecolor="white"); plt.close(fig)

    envs = ["Shop", "Coding", "DeepResearch", "AutoResearch"]
    camg = [97.1, 26.6, 55.5, 38.3]; compact = [97.5, 19.5, 50.8, 24.2]
    x = np.arange(len(envs)); width = .35
    fig, ax = plt.subplots(figsize=(8.2, 3.2), dpi=220)
    ax.bar(x-width/2, camg, width, label="CAMG-RL", color="#2B628F")
    ax.bar(x+width/2, compact, width, label="CompactionRL", color="#AFC2D1")
    ax.set_ylim(0, 110); ax.set_ylabel("Success rate (%)"); ax.set_xticks(x, envs)
    ax.spines[["top", "right"]].set_visible(False); ax.grid(axis="y", color="#E6E8EA", linewidth=.7); ax.set_axisbelow(True)
    ax.legend(frameon=False, ncol=2, loc="upper left", bbox_to_anchor=(0, 1.13), fontsize=9)
    for vals, dx in [(camg, -width/2), (compact, width/2)]:
        for xx, v in zip(x, vals): ax.text(xx+dx, v+2, f"{v:.1f}", ha="center", fontsize=8)
    fig.tight_layout(); fig.savefig(out_dir / "environment-result.png", facecolor="white"); plt.close(fig)

    models = ["4B", "9B", "27B", "35B-A3B", "122B-A10B", "397B-A17B", "CAMG-RL-4B", "CAMG-RL-9B"]
    swe = [7.6, 14.8, 17.4, 15.6, 22.0, 34.4, 15.8, 27.6]
    mle = [0.0, 4.5, 4.5, 4.5, 9.1, 13.6, 4.5, 9.1]
    x = np.arange(len(models)); width = .36
    fig, ax = plt.subplots(figsize=(9.0, 3.3), dpi=220)
    ax.bar(x-width/2, swe, width, label="SWE-bench Verified", color="#2B628F")
    ax.bar(x+width/2, mle, width, label="MLE-bench Lite", color="#D7A2A8")
    ax.set_ylim(0, 40); ax.set_ylabel("Success rate (%)"); ax.set_xticks(x, models, rotation=25, ha="right")
    ax.spines[["top", "right"]].set_visible(False); ax.grid(axis="y", color="#E6E8EA", linewidth=.7); ax.set_axisbelow(True)
    ax.legend(frameon=False, ncol=2, loc="upper left", fontsize=9)
    for i in [6, 7]:
        ax.axvspan(i-.5, i+.5, color="#F8E5E8", zorder=-1)
    fig.tight_layout(); fig.savefig(out_dir / "transfer-result.png", facecolor="white"); plt.close(fig)


def build(out_pptx: Path, site_root: Path):
    assets = site_root / "assets"
    work = out_pptx.parent / "generated"
    make_chart_images(work)
    prs = Presentation(); prs.slide_width = Inches(SW); prs.slide_height = Inches(SH)
    blank = prs.slide_layouts[6]

    # 1. Cover
    s = prs.slides.add_slide(blank)
    text(s, "CAMG / CAMG-RL", 0.60, 0.48, 5.4, 0.22, size=10, color=RED, bold=True)
    text(s, "JD.com", 8.50, .48, .9, .25, size=12, bold=True, align=PP_ALIGN.RIGHT)
    text(s, "Coding Agent Memory\nPost-training", 0.60, 1.02, 8.8, 1.15, size=32, bold=True)
    text(s, "Unlocking the Memory Potential of Pre-trained File Operations\nfor Long-Horizon Tasks via Reinforcement Learning", 0.62, 2.42, 8.7, .66, size=18, color=MID)
    text(s, "Lirui Luo · Kelong Mao · Heming Xia · Rongqing Li · Xinwei Yang · Luyu Chen\nKieran Wong · Yudong Guo · Xinrui Wang · Jiayin Zhu · Simiu Gu · Sulong Xu · Cong Fang", 0.62, 3.54, 8.7, .60, size=11, color=MID)
    text(s, "JD.com", .62, 4.24, 1.3, .23, size=11, bold=True)
    for label, url, y in [("Paper  ·  arXiv:2609.34422", "https://arxiv.org/abs/2609.34422", 4.85),
                          ("Project  ·  liruiluo.github.io/agentmemorygym", "https://liruiluo.github.io/agentmemorygym/", 5.13)]:
        box = text(s, label, .62, y, 7.8, .22, size=10, color=BLUE)
        box.text_frame.paragraphs[0].runs[0].hyperlink.address = url

    # 2. Motivation
    s = prs.slides.add_slide(blank); title(s, "Long tasks need state that outlives context", 2)
    image(s, assets / "teaser.png", .55, .98, 8.9, 3.65)
    text(s, "Useful state survives a context replacement.", .80, 4.62, 8.4, .28, size=16, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    caption(s, "MOTIVATION", "A bounded active context is not a durable workspace.")

    # 3. Question
    s = prs.slides.add_slide(blank); title(s, "Can RL unlock the memory prior in file operations?", 3)
    rect(s, .62, 1.18, 3.75, 2.35, fill=BLUE_LIGHT, line=BLUE)
    text(s, "PRE-TRAINING PRIOR", .88, 1.46, 2.9, .23, size=10, color=BLUE, bold=True)
    text(s, "Shell commands, files, search,\nand revision are familiar actions.", .88, 1.90, 3.12, 1.25, size=22, bold=True)
    rect(s, 5.02, 1.18, 4.32, 2.35, fill=RED_LIGHT, line=RED)
    text(s, "OPEN QUESTION", 5.30, 1.46, 2.2, .23, size=10, color=RED, bold=True)
    text(s, "Can downstream task reward\nturn that prior into long-horizon memory?", 5.30, 1.86, 3.45, 1.30, size=22, bold=True)
    line(s, 4.38, 2.33, 4.98, 2.33, color=RED, width=2)
    text(s, "→", 4.56, 2.10, .25, .36, size=23, color=RED, bold=True, align=PP_ALIGN.CENTER)
    caption(s, "RESEARCH QUESTION", "Learn memory behavior through the task objective, without inventing a new memory tool vocabulary.")

    # 4. Contributions
    s = prs.slides.add_slide(blank); title(s, "CAMG spans four native task worlds", 4)
    image(s, assets / "task-examples.png", .55, .94, 8.9, 3.25)
    bullets(s, ["four environments: Shop, Coding, DeepResearch, and AutoResearch", "native task interface + shell + episode-persistent workspace", "one shared policy and downstream reward path"], .78, 4.30, 8.2, .63, size=12, gap=3)
    caption(s, "CAMG", "The Gym holds task semantics fixed while exposing a common file-backed memory substrate.")

    # 5. Architecture
    s = prs.slides.add_slide(blank); title(s, "One policy learns across all four environments", 5)
    image(s, assets / "framework.png", .58, .92, 8.85, 3.72)
    text(s, "native task action", .75, 4.66, 2.1, .25, size=12, color=BLUE, bold=True)
    text(s, "+", 2.82, 4.66, .24, .25, size=16, color=RED, bold=True, align=PP_ALIGN.CENTER)
    text(s, "filesystem action", 3.18, 4.66, 2.1, .25, size=12, color=BLUE, bold=True)
    text(s, "→", 5.32, 4.66, .24, .25, size=16, color=RED, bold=True, align=PP_ALIGN.CENTER)
    text(s, "shared task reward", 5.70, 4.66, 2.0, .25, size=12, color=BLUE, bold=True)
    caption(s, "SYSTEM", "Fully asynchronous training absorbs variable episode durations while keeping a single actor and critic.")

    # 6. Method
    s = prs.slides.add_slide(blank); title(s, "Files carry reasoning across context boundaries", 6)
    rect(s, .68, 1.14, 2.55, 2.76, fill=WHITE, line=BLUE, width=1.3)
    rect(s, 3.74, 1.14, 2.55, 2.76, fill=WHITE, line=RED, width=1.3)
    rect(s, 6.80, 1.14, 2.55, 2.76, fill=WHITE, line=GREEN, width=1.3)
    for x, tag, head, body, col in [(.95,"ACT","Work on the task","Use native tools and\nordinary shell actions.",BLUE),(4.01,"WRITE","Save what matters","Record evidence, plans,\nand the next step.",RED),(7.07,"RETURN","Read it back","After replacement,\nretrieve the state.",GREEN)]:
        text(s, tag, x, 1.43, 1.7, .24, size=10, color=col, bold=True)
        text(s, head, x, 1.91, 1.9, .52, size=18, bold=True)
        text(s, body, x, 2.62, 1.92, .62, size=14, color=MID)
    text(s, "→", 3.37, 2.37, .25, .35, size=22, color=RED, bold=True)
    text(s, "→", 6.43, 2.37, .25, .35, size=22, color=RED, bold=True)
    rect(s, 2.84, 4.24, 4.34, .52, fill=BLUE_LIGHT, line=BLUE, width=1.0)
    text(s, ".agent_memory/CONTINUATION.md", 2.84, 4.24, 4.34, .52, size=15, color=BLUE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    caption(s, "MEMORY AS REASONING", "No separate hidden memory channel: editable files carry plans, evidence, and intermediate findings.")

    # 7. Learning signal
    s = prs.slides.add_slide(blank); title(s, "Task reward trains the whole memory chain", 7)
    text(s, "One task reward", .74, 1.08, 3.0, .35, size=22, bold=True)
    text(s, "One shared credit path", 5.85, 1.08, 3.4, .35, size=22, bold=True)
    line(s, 1.02, 2.03, 8.96, 2.03, color=LINE, width=1.3)
    for x, lab, sub, col in [(1.00,"write","make state durable",BLUE),(3.34,"retrieve","bring state back",RED),(5.70,"use","change the next decision",GREEN),(8.03,"reward","succeed or fail",INK)]:
        rect(s, x, 1.70, 1.45, .68, fill=WHITE, line=col, width=1.25)
        text(s, lab, x, 1.81, 1.45, .25, size=14, color=col, bold=True, align=PP_ALIGN.CENTER)
        text(s, sub, x-.12, 2.48, 1.7, .35, size=11, color=MID, align=PP_ALIGN.CENTER)
        if x < 8: text(s, "→", x+1.57, 1.86, .22, .28, size=18, color=RED, bold=True, align=PP_ALIGN.CENTER)
    rect(s, 1.03, 3.52, 7.90, .92, fill=RED_LIGHT, line=RED, width=1.2)
    text(s, "Native task reward provides the learning signal.", 1.03, 3.72, 7.90, .34, size=21, color=RED, bold=True, align=PP_ALIGN.CENTER)
    caption(s, "CREDIT ASSIGNMENT", "Task and memory-file responses share task returns. The mechanical context reset creates no policy action.")

    # 8. Main result
    s = prs.slides.add_slide(blank); title(s, "CAMG-RL achieves 54.4% average success", 8)
    image(s, work / "main-result.png", .62, 1.05, 8.72, 3.32)
    rich_text(s, [{"text": "+6.4 percentage points", "color": RED, "bold": True},
                  {"text": " over CompactionRL"}], .85, 4.54, 8.3, .35, size=20, align=PP_ALIGN.CENTER)
    caption(s, "TABLE 2", "128 test tasks per environment; matched decoding, action budgets, runtimes, and graders.")

    # 9. By environment
    s = prs.slides.add_slide(blank); title(s, "The largest gains are in Coding and AutoResearch", 9)
    image(s, work / "environment-result.png", .72, 1.03, 8.56, 3.17)
    bullets(s, ["Shop: both methods nearly saturate the task", "Coding: 26.6% vs. 19.5% for learned compaction", "AutoResearch: 38.3% vs. 24.2% — the largest absolute gap"], .86, 4.31, 8.15, .64, size=12, gap=3)
    caption(s, "RESULT", "A single policy learns one memory behavior while each environment keeps its native task semantics.")

    # 10. Interface prior
    s = prs.slides.add_slide(blank); title(s, "File operations build on a measurable pre-training prior", 10)
    image(s, assets / "interface-prior.png", .58, 1.04, 8.9, 3.75)
    text(s, "code pre-training supplies the interface prior", .85, 4.72, 8.0, .25, size=16, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    caption(s, "INTERFACE PRIOR", "Training mainly has to discover when to write, revise, retrieve, and reuse — not a new action language.")

    # 11. External transfer
    s = prs.slides.add_slide(blank); title(s, "Learned memory transfers beyond the training Gym", 11)
    image(s, work / "transfer-result.png", .50, 1.02, 9.05, 3.62)
    rect(s, .85, 4.63, 3.86, .43, fill=BLUE_LIGHT, line=BLUE, width=1.0)
    text(s, "CAMG-RL-4B: 15.8 / 4.5", .85, 4.72, 3.86, .18, size=11.5, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    rect(s, 5.28, 4.63, 3.86, .43, fill=RED_LIGHT, line=RED, width=1.0)
    text(s, "CAMG-RL-9B: 27.6 / 9.1", 5.28, 4.72, 3.86, .18, size=11.5, color=RED, bold=True, align=PP_ALIGN.CENTER)
    caption(s, "EXTERNAL TRANSFER", "The 4B and 9B policies are competitive with much larger frozen Qwen3.5 models on the paired benchmarks.")

    # 12. Ablation
    s = prs.slides.add_slide(blank); title(s, "Memory-action credit and general files both matter", 12)
    image(s, assets / "ablation.png", .58, 1.00, 8.86, 3.70)
    text(s, "w/o memory-action gradients", .88, 4.68, 2.45, .24, size=12, color=MID, bold=True, align=PP_ALIGN.CENTER)
    text(s, "continuation file only", 3.85, 4.68, 2.10, .24, size=12, color=MID, bold=True, align=PP_ALIGN.CENTER)
    text(s, "full CAMG-RL", 7.05, 4.68, 1.64, .24, size=12, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    caption(s, "FIGURE 4", "Average success: 54.4% full; 21.1% without memory-action gradients; 41.8% when retrained with continuation-only memory.")

    # 13. Memory-content intervention
    s = prs.slides.add_slide(blank); title(s, "Changing saved memory content reduces later success", 13)
    text(s, "Intervene at the first context replacement", .75, 1.14, 8.5, .4, size=21, bold=True)
    text(s, "Continue the trajectory after replacing the saved files.", .75, 1.68, 8.5, .30, size=16, color=MID)
    for y, label, value, width in [(2.53, "Blank memory", "−14.1 pts", 4.50),
                                  (3.45, "Task-mismatched memory", "−11.5 pts", 3.67)]:
        text(s, label, .75, y+.06, 3.15, .34, size=16, bold=True)
        rect(s, 4.12, y, width, .48, fill=RED_LIGHT, line=RED, width=.7, rounded=False)
        text(s, value, 4.27, y+.08, width-.3, .30, size=18, color=RED, bold=True)
    text(s, "Average success-rate change from the matched original-memory condition", .80, 4.50, 8.4, .28, size=13, color=MID)
    caption(s, "APPENDIX D.3 · FIGURE 8", "The largest drop occurs in Shop. Saved content affects subsequent task decisions.")

    # 14. Case
    s = prs.slides.add_slide(blank); title(s, "A coding episode makes the memory chain visible", 14)
    image(s, assets / "coding-case.png", .50, .97, 9.04, 3.90)
    caption(s, "FIGURE 5 · TRAINING CASE", "Reproduce, patch, save before context replacement, read back, verify, and submit.")

    # 15. Takeaways
    s = prs.slides.add_slide(blank); title(s, "Takeaways", 15)
    text(s, "01", .70, 1.02, .54, .34, size=16, color=RED, bold=True)
    text(s, "CAMG turns long-horizon memory into a testable, multi-domain RL environment suite.", 1.44, 1.00, 7.55, .58, size=18, bold=True)
    line(s, .72, 1.67, 9.22, 1.67, color=LINE, width=.8)
    text(s, "02", .70, 2.05, .54, .34, size=16, color=RED, bold=True)
    text(s, "CAMG-RL learns to use ordinary files as memory through downstream task reward.", 1.44, 2.03, 7.55, .58, size=18, bold=True)
    line(s, .72, 2.70, 9.22, 2.70, color=LINE, width=.8)
    text(s, "03", .70, 3.08, .54, .34, size=16, color=RED, bold=True)
    text(s, "A small post-trained policy can transfer this behavior to external agent benchmarks.", 1.44, 3.06, 7.55, .58, size=18, bold=True)
    rect(s, .72, 4.20, 8.50, .72, fill=BLUE_LIGHT, line=BLUE, width=1.0)
    text(s, "Files are not just artifacts. They are memory.", .72, 4.36, 8.50, .42, size=21, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    text(s, "Coding Agent Memory Post-training · arXiv:2609.34422", .72, 5.12, 5.0, .18, size=8.5, color=LIGHT)

    out_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.core_properties.title = "Coding Agent Memory Post-training"
    prs.core_properties.subject = "Public paper talk · arXiv:2609.34422v1"
    prs.core_properties.author = "Lirui Luo et al. · JD.com"
    prs.save(out_pptx)
    print(out_pptx)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--site-root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    build(args.out.resolve(), args.site_root.resolve())


if __name__ == "__main__":
    main()
