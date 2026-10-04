#!/usr/bin/env python3
"""Render this deck's native text, images and geometry without a GUI."""
from __future__ import annotations

import argparse
import io
import json
import math
from html import escape
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph


def pt(value):
    return value / 12700


def color(c, fallback="111111"):
    try:
        return str(c.rgb)
    except (AttributeError, TypeError):
        return fallback


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", type=Path)
    ap.add_argument("pdf", type=Path)
    args = ap.parse_args()
    font_root = Path("/System/Library/Fonts/Supplemental")
    for name, fname in [("Arial", "Arial.ttf"), ("Arial-Bold", "Arial Bold.ttf"),
                        ("Arial-Italic", "Arial Italic.ttf"),
                        ("Arial-BoldItalic", "Arial Bold Italic.ttf")]:
        pdfmetrics.registerFont(TTFont(name, str(font_root / fname)))
    for index, suffix in enumerate(["", "-Bold", "-Italic", "-BoldItalic"]):
        pdfmetrics.registerFont(TTFont("Menlo" + suffix,
                                     "/System/Library/Fonts/Menlo.ttc", subfontIndex=index))
    prs = Presentation(args.pptx)
    sw, sh = pt(prs.slide_width), pt(prs.slide_height)
    args.pdf.parent.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(args.pdf), pagesize=(sw, sh))
    pdf.setTitle(prs.core_properties.title)
    pdf.setAuthor(prs.core_properties.author)
    overflow = []
    for index, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            x, y, w, h = map(pt, (shape.left, shape.top, shape.width, shape.height))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pdf.drawImage(ImageReader(io.BytesIO(shape.image.blob)), x, sh-y-h, w, h, mask="auto")
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.LINE:
                pdf.setStrokeColor(HexColor("#" + color(shape.line.color)))
                pdf.setFillColor(HexColor("#" + color(shape.line.color)))
                pdf.setLineWidth(shape.line.width.pt if shape.line.width else 1)
                transform = shape._element.xpath(".//a:xfrm")[0]
                flip_h = transform.get("flipH") in ("1", "true")
                flip_v = transform.get("flipV") in ("1", "true")
                x1, x2 = (x+w, x) if flip_h else (x, x+w)
                y1, y2 = (sh-y-h, sh-y) if flip_v else (sh-y, sh-y-h)
                if shape._element.xpath(".//a:prstDash"):
                    pdf.setDash(4, 3)
                pdf.line(x1, y1, x2, y2)
                pdf.setDash()
                for end, tip, start in [("tailEnd", (x2, y2), (x1, y1)),
                                         ("headEnd", (x1, y1), (x2, y2))]:
                    if not shape._element.xpath(f'.//a:{end}[@type="triangle"]'):
                        continue
                    dx, dy = tip[0]-start[0], tip[1]-start[1]
                    length = math.hypot(dx, dy)
                    if not length:
                        continue
                    ux, uy = dx/length, dy/length
                    extent = max(5, pdf._lineWidth*3)
                    path = pdf.beginPath()
                    path.moveTo(*tip)
                    path.lineTo(tip[0]-extent*ux+extent*.45*uy,
                                tip[1]-extent*uy-extent*.45*ux)
                    path.lineTo(tip[0]-extent*ux-extent*.45*uy,
                                tip[1]-extent*uy+extent*.45*ux)
                    path.close()
                    pdf.drawPath(path, stroke=0, fill=1)
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                pdf.setFillColor(HexColor("#" + color(shape.fill.fore_color, "FFFFFF")))
                pdf.setStrokeColor(HexColor("#" + color(shape.line.color)))
                pdf.setLineWidth(shape.line.width.pt if shape.line.width else 1)
                if "ROUNDED" in str(shape.auto_shape_type):
                    pdf.roundRect(x, sh-y-h, w, h, min(w, h)*.10, fill=1, stroke=1)
                elif "OVAL" in str(shape.auto_shape_type):
                    pdf.ellipse(x, sh-y-h, x+w, sh-y, fill=1, stroke=1)
                else:
                    pdf.rect(x, sh-y-h, w, h, fill=1, stroke=1)
            if not shape.has_text_frame or not shape.text:
                continue
            tf = shape.text_frame
            ml, mt, mr, mb = map(pt, (tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom))
            width = w-ml-mr
            paragraphs = []
            for p in tf.paragraphs:
                if not p.runs:
                    continue
                parts = []
                max_size = max((r.font.size.pt if r.font.size else 18) for r in p.runs)
                for r in p.runs:
                    family = "Menlo" if r.font.name == "Menlo" else "Arial"
                    if r.font.bold and r.font.italic:
                        family += "-BoldItalic"
                    elif r.font.bold:
                        family += "-Bold"
                    elif r.font.italic:
                        family += "-Italic"
                    size = r.font.size.pt if r.font.size else max_size
                    content = escape(r.text)
                    if r.font.name == "Menlo":
                        content = content.replace(" ", "&#160;")
                    part = '<font name="{}" size="{}" color="#{}">{}</font>'.format(family, size, color(r.font.color), content)
                    if r.hyperlink.address:
                        part = '<link href="{}">{}</link>'.format(escape(r.hyperlink.address, quote=True), part)
                    parts.append(part)
                leading = max_size * (p.line_spacing if isinstance(p.line_spacing, float) else 1.0)
                # PowerPoint alignment values start at one; ReportLab at zero.
                align = max(0, min(2, int(p.alignment or 1)-1))
                style = ParagraphStyle("slide", fontName="Arial", fontSize=max_size,
                                       leading=leading, alignment=align,
                                       textColor=HexColor("#111111"),
                                       spaceAfter=p.space_after.pt if p.space_after else 0)
                para = Paragraph("".join(parts), style)
                _, height = para.wrap(width, 10000)
                paragraphs.append((para, height, style.spaceAfter))
            height_total = sum(hh + gap for _, hh, gap in paragraphs)
            if paragraphs:
                height_total -= paragraphs[-1][2]
            available = h-mt-mb
            if height_total > available + .8:
                overflow.append({"slide": index, "text": shape.text[:90], "height": round(height_total, 1), "available": round(available, 1)})
            yy = sh-y-mt
            if tf.vertical_anchor == MSO_ANCHOR.MIDDLE:
                yy -= (available-height_total)/2
            elif tf.vertical_anchor == MSO_ANCHOR.BOTTOM:
                yy -= available-height_total
            for para, height, gap in paragraphs:
                para.drawOn(pdf, x+ml, yy-height)
                yy -= height+gap
        pdf.showPage()
    pdf.save()
    print(json.dumps({"pages":len(prs.slides),"pdf":str(args.pdf),"overflow":overflow}, ensure_ascii=False))
    if overflow:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
